import json, base64, io, os, traceback
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BOOK2 = os.path.join(os.path.dirname(__file__), 'Book2.xlsx')

def handler(event, context):
    try:
        body = event.get("body", "{}")
        if event.get("isBase64Encoded"):
            body = base64.b64decode(body).decode("utf-8")
        data = json.loads(body)
        result = generate(data)
        return {"statusCode":200,"headers":{"Content-Type":"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet","Access-Control-Allow-Origin":"*","Content-Disposition":"attachment; filename=comparison.xlsx"},"body":result,"isBase64Encoded":True}
    except Exception as e:
        return {"statusCode":500,"headers":{"Content-Type":"application/json","Access-Control-Allow-Origin":"*"},"body":json.dumps({"error":str(e),"trace":traceback.format_exc()})}

def generate(data):
    wb_ref = load_workbook(BOOK2)
    ws_ref = wb_ref['Landscape R01']
    ws_sum = wb_ref['Summary ']

    def ap(dst, addr, s='c'):
        src = ws_sum[addr] if s=='s' else ws_ref[addr]
        f=src.font
        if f: dst.font=Font(name=f.name or'Arial',bold=f.bold,italic=f.italic,size=f.size or 14,color=f.color.rgb if f.color and f.color.type=='rgb' else'000000')
        b=src.border
        if b:
            def cs(x): return Side(style=x.style,color=x.color.rgb if x.color and x.color.type=='rgb' else'000000') if x and x.style else Side(style=None)
            dst.border=Border(left=cs(b.left),right=cs(b.right),top=cs(b.top),bottom=cs(b.bottom))
        a=src.alignment
        if a: dst.alignment=Alignment(horizontal=a.horizontal,vertical=a.vertical,wrap_text=a.wrap_text)

    def W(ws,r,c,v,sa=None,s='c',nf=None):
        cell=ws.cell(row=r,column=c,value=v)
        if sa: ap(cell,sa,s)
        if nf: cell.number_format=nf
        return cell

    def MG(ws,r1,c1,r2,c2):
        ws.merge_cells(start_row=r1,start_column=c1,end_row=r2,end_column=c2)

    NUM="#,##0.00"; PCT="0.0%"; QTY="#,##0.##"
    wb=Workbook(); wb.remove(wb.active)

    pkg=data.get("pkg",{}); proj=data.get("proj",{})
    scData=data.get("scData",[]); rows=data.get("boqRows",[]); lowMap=data.get("lowestPriceMap",{})

    if not data.get("summaryOnly"):
        ws=wb.create_sheet("Comparison Statement")
        # Row heights from Book2
        for r,d in ws_ref.row_dimensions.items():
            if d.height: ws.row_dimensions[r].height=d.height
        # Col widths
        for l,d in ws_ref.column_dimensions.items():
            if d.width: ws.column_dimensions[l].width=d.width

        # SC column layout: 3 cols per rev (Price|Qty|Amount) matching Book2
        SC=8; sc_cols={}; col=SC
        for d in scData:
            sc_cols[d["sc"]["id"]]=col; col+=max(1,len(d["revs"]))*3
        LAST=col-1

        # Update SC col widths
        for d in scData:
            sc_col=sc_cols[d["sc"]["id"]]
            for ri in range(max(1,len(d["revs"]))):
                ws.column_dimensions[get_column_letter(sc_col+ri*3)].width=19.82
                ws.column_dimensions[get_column_letter(sc_col+ri*3+1)].width=18.45
                ws.column_dimensions[get_column_letter(sc_col+ri*3+2)].width=23.36

        # Row 2: Title
        c2=ws.cell(row=2,column=2,value=f"COMPARISON STATEMENT FOR {pkg.get('name','').upper()} (SUPPLY & INSTALLATION)")
        ap(c2,'E2'); MG(ws,2,2,2,LAST)

        # Row 4: Project
        W(ws,4,3,"Project Name:",'C4')
        W(ws,4,4,f"{proj.get('name','')} {proj.get('loc','')}".strip(),'D4')
        MG(ws,4,4,4,LAST-4)
        W(ws,4,LAST-3,"DOC Ref:",'AY4'); W(ws,4,LAST-2,pkg.get("ref",""),'AZ4')
        W(ws,4,LAST-1,"DATE",'BJ4'); W(ws,4,LAST,pkg.get("rev","REV-0"),'BL4')

        # Rows 5-6: Headers + SC names/contacts (merged)
        W(ws,5,2,"Sno",'B5'); W(ws,5,3,"Description",'C5'); W(ws,5,4,"Unit",'D5')
        W(ws,6,5,"Quantity",'E6'); W(ws,6,6,"BCC Rate",'F6'); W(ws,6,7,"BCC Budget",'G6')
        MG(ws,5,2,6,2); MG(ws,5,3,6,3); MG(ws,5,4,6,4)
        for d in scData:
            sc_col=sc_cols[d["sc"]["id"]]; sp=max(1,len(d["revs"]))*3
            W(ws,5,sc_col,d["sc"]["company"],'H5')
            if sp>1: MG(ws,5,sc_col,5,sc_col+sp-1)
            ct=" ".join(filter(None,[d["sc"].get("contact",""),d["sc"].get("mobile","")]))
            W(ws,6,sc_col,ct,'H6')
            if sp>1: MG(ws,6,sc_col,6,sc_col+sp-1)

        # Rows 7-9: Ref/Date/Type per SC
        for row_n,key,sa in [(7,"qref",'H7'),(8,"qdate",'H8'),(9,"ctype",'H9')]:
            for d in scData:
                sc_col=sc_cols[d["sc"]["id"]]; sp=max(1,len(d["revs"]))*3
                W(ws,row_n,sc_col,d["info0"].get(key,"Lumpsum" if key=="ctype" else ""),sa)
                if sp>1: MG(ws,row_n,sc_col,row_n,sc_col+sp-1)

        # Row 11: Rev headers
        W(ws,11,6,pkg.get("rev","REV-0"),'F6'); W(ws,11,7,"Total",'G6')
        for d in scData:
            sc_col=sc_cols[d["sc"]["id"]]; n=max(1,len(d["revs"]))
            for ri in range(n):
                W(ws,11,sc_col+ri*3,  f"R{ri} Price", 'H11')
                W(ws,11,sc_col+ri*3+1,"Qty",          'I11')
                W(ws,11,sc_col+ri*3+2,f"R{ri} Amount",'J11')

        # Data rows from 13
        dRow=13; sno=0; bud=0.0
        for row in rows:
            isH=row.get("isHeader") or row.get("type") in ("header","section")
            if isH:
                W(ws,dRow,2,"",'B18'); W(ws,dRow,3,row.get("description",""),'C12')
                MG(ws,dRow,3,dRow,LAST); dRow+=1; continue
            sno+=1
            qty=float(row["quantity"]) if row.get("quantity") else None
            rate=float(row["budgetRate"]) if row.get("budgetRate") else None
            isX=row.get("_scAdded",False)
            W(ws,dRow,2,sno if qty else "",'B18',nf='0' if qty else None)
            W(ws,dRow,3,row.get("description","")+((" ★") if isX else ""),'C18')
            W(ws,dRow,4,row.get("unit",""),'D18')
            W(ws,dRow,5,qty,'E18',nf=QTY)
            if rate and qty:
                bud+=rate*qty
                W(ws,dRow,6,rate,'F18',nf=NUM); W(ws,dRow,7,rate*qty,'G18',nf=NUM)
            else:
                W(ws,dRow,6,None,'F18'); W(ws,dRow,7,None,'G18')
            for d in scData:
                sc_col=sc_cols[d["sc"]["id"]]; ab=row.get("_addedByScId"); n=max(1,len(d["revs"]))
                for ri in range(n):
                    rv=d["revs"][ri] if ri<len(d["revs"]) else None
                    cp=sc_col+ri*3; cq=cp+1; ca=cp+2
                    if isX and d["sc"]["id"]!=ab:
                        low=lowMap.get(row.get("id"),0) or 0
                        c=ws.cell(row=dRow,column=cp,value="EXCL"); ap(c,'H18')
                        c.font=Font(name='Arial',bold=True,size=14,color='FF0000')
                        ws.cell(row=dRow,column=cq); ap(ws.cell(row=dRow,column=cq),'I18')
                        c=ws.cell(row=dRow,column=ca,value=float(low)); ap(c,'J18')
                        c.font=Font(name='Arial',size=14,color='FF0000'); c.number_format=NUM
                        continue
                    item=next((x for x in(rv["items"] if rv else[]) if x.get("boqId")==row.get("id")),None)
                    if not item:
                        W(ws,dRow,cp,None,'H18'); W(ws,dRow,cq,None,'I18'); W(ws,dRow,ca,None,'J18'); continue
                    opt=item.get("option","priced")
                    if opt=="excluded":
                        low=lowMap.get(row.get("id"),0) or 0
                        c=ws.cell(row=dRow,column=cp,value="EXCL"); ap(c,'H18')
                        c.font=Font(name='Arial',bold=True,size=14,color='FF0000')
                        ap(ws.cell(row=dRow,column=cq),'I18')
                        c=ws.cell(row=dRow,column=ca,value=float(low)); ap(c,'J18')
                        c.font=Font(name='Arial',size=14,color='FF0000'); c.number_format=NUM
                    elif opt=="included":
                        W(ws,dRow,cp,"INCLUDED",'H18'); W(ws,dRow,cq,None,'I18'); W(ws,dRow,ca,"INCLUDED",'J18')
                    elif opt=="rateonly":
                        r2=float(item.get("rate") or 0)
                        W(ws,dRow,cp,r2 or None,'H18',nf=NUM); W(ws,dRow,cq,None,'I18'); W(ws,dRow,ca,"Rate Only",'J18')
                    else:
                        r2=float(item.get("rate") or 0); q2=float(item.get("qty") or qty or 0)
                        W(ws,dRow,cp,r2 or None,'H18',nf=NUM)
                        W(ws,dRow,cq,q2 or qty,'I18',nf=QTY)
                        W(ws,dRow,ca,r2*q2 if r2 and q2 else None,'J18',nf=NUM)
            dRow+=1

        # Total row
        TR=dRow
        W(ws,TR,2,"",'B130'); W(ws,TR,3,"TOTAL AMOUNT",'C130'); MG(ws,TR,3,TR,5)
        W(ws,TR,6,"",'F18'); W(ws,TR,7,bud,'G130',nf=NUM)
        for d in scData:
            sc_col=sc_cols[d["sc"]["id"]]; n=max(1,len(d["revs"]))
            for ri in range(n):
                tot=d["revTotals"][ri] if ri<len(d["revTotals"]) else 0
                W(ws,TR,sc_col+ri*3,"",'H130'); W(ws,TR,sc_col+ri*3+1,"",'I18')
                W(ws,TR,sc_col+ri*3+2,float(tot),'J130',nf=NUM)
        dRow+=1

        # Savings/Loss row
        SR=dRow
        W(ws,SR,2,"",'B130'); W(ws,SR,3,"Savings / (Loss) vs Budget (%)",'C130'); MG(ws,SR,3,SR,7)
        for d in scData:
            sc_col=sc_cols[d["sc"]["id"]]; n=max(1,len(d["revs"]))
            for ri in range(n):
                tot=d["revTotals"][ri] if ri<len(d["revTotals"]) else 0
                pct=(bud-tot)/bud if bud and tot else 0
                W(ws,SR,sc_col+ri*3,"",'H130'); W(ws,SR,sc_col+ri*3+1,"",'I18')
                c=W(ws,SR,sc_col+ri*3+2,pct,'J130',nf=PCT)
                c.font=Font(name='Arial',bold=True,size=14,color='FF0000' if pct<0 else '00B050')
        dRow+=2

        # Terms
        W(ws,dRow,3,"Specified Manufacturer",'C130'); MG(ws,dRow,3,dRow,7)
        for d in scData:
            sc_col=sc_cols[d["sc"]["id"]]; sp=max(1,len(d["revs"]))*3
            W(ws,dRow,sc_col,d["terms"].get("mat",""),'C18')
            if sp>1: MG(ws,dRow,sc_col,dRow,sc_col+sp-1)
        dRow+=1
        W(ws,dRow,3,"COMMERCIAL TERMS & CONDITIONS",'C130'); MG(ws,dRow,3,dRow,LAST); dRow+=1
        for lbl,key in [("Advance (%)","adv"),("Advance Payment Guarantee (SC/BG)","advg"),
                         ("Performance (10%)","perf"),("Performance Guarantee (SC/BG)","perfg"),
                         ("Retention 10% (5% TOC / 5% DLP)","ret"),("Material on Site %","mos"),
                         ("Progress Payment %","prog"),("Payment Terms","pay"),
                         ("Delivery Conditions","del"),("Warranty","warr"),
                         ("Proposed Material","mat"),("Remarks / Exclusions","notes")]:
            W(ws,dRow,3,lbl,'C18'); MG(ws,dRow,3,dRow,7)
            for d in scData:
                sc_col=sc_cols[d["sc"]["id"]]; sp=max(1,len(d["revs"]))*3
                W(ws,dRow,sc_col,d["terms"].get(key,""),'C18')
                if sp>1: MG(ws,dRow,sc_col,dRow,sc_col+sp-1)
            dRow+=1
        dRow+=1
        for i,(nm,ttl) in enumerate([("Name: Ms. Reshma Ashok","Sr. Procurement Engineer"),
                                       ("Name: Mr. Makesh Nagarajan","HOD - Procurement"),
                                       ("Name: Mr. Asad Shaikh","Cost Control Manager"),
                                       ("Name: Mr. Danny Kurian","Director - BCC")]):
            c=SC+i*3
            W(ws,dRow,c,nm,'C18'); W(ws,dRow+1,c,ttl,'C18')
            W(ws,dRow+2,c,"Signature: _______________________",'C18')
            W(ws,dRow+3,c,"Date: ____________________________",'C18')
        ws.freeze_panes="C13"; ws.print_title_rows="2:11"

    # Summary sheet
    ws2=wb.create_sheet("Summary")
    for r,d in ws_sum.row_dimensions.items():
        if d.height: ws2.row_dimensions[r].height=d.height
    for l,d in ws_sum.column_dimensions.items():
        if d.width: ws2.column_dimensions[l].width=d.width
    proj=data.get("proj",{}); allPkgs=data.get("allPkgs",[data.get("pkg",{})])
    for i,(lbl,val) in enumerate([("Project:",proj.get("name","")),("Project Code:",proj.get("code","")),
                                    ("Client Name:",proj.get("client","")),("Consultant:",proj.get("cons","")),
                                    ("Package Name:",allPkgs[0].get("name","") if allPkgs else "")],start=1):
        W(ws2,i,1,lbl,f'A{i}',s='s'); W(ws2,i,3,val,f'C{i}',s='s')
        ws2.merge_cells(start_row=i,start_column=3,end_row=i,end_column=16)
    W(ws2,6,1,"Bid Summary",'A6',s='s')
    for addr in ['A8','B8','C8','D8','J8','N8','P8','Q8','D9','E9','F9','G9','H9','I9','J9','K9','L9','M9','N9','O9']:
        src=ws_sum[addr]; dst=ws2.cell(row=src.row,column=src.column,value=src.value); ap(dst,addr,'s')
    for m in ws_sum.merged_cells.ranges:
        try: ws2.merge_cells(str(m))
        except: pass
    row=10; gb=0
    for pkg2 in allPkgs:
        pkgSubs=pkg2.get("submissions",{}); pkgScs=pkg2.get("subcontractors",[])
        budget=sum(float(r.get("quantity") or 0)*float(r.get("budgetRate") or 0)
                   for r in pkg2.get("boqRows",[]) if not r.get("isHeader") and r.get("quantity") and r.get("budgetRate"))
        gb+=budget; sc_rows=[]
        for sc in pkgScs:
            sub=pkgSubs.get(sc["id"],{}); revs=sub.get("revHistory",[])
            if not revs and sub.get("rev0"): revs=[{"items":sub["rev0"],"info":sub.get("r0info",{}),"terms":sub.get("r0terms",{})}]
            if not revs: continue
            rt=[sum(float(x.get("amount") or 0) for x in rv["items"] if isinstance(x.get("amount"),(int,float)) and x.get("option")!="rateonly") for rv in revs]
            sc_rows.append({"sc":sc,"rt":rt,"fo":rt[-1] if rt else 0,"ft":revs[-1].get("terms",{}),"info0":revs[0].get("info",{})})
        sc_rows.sort(key=lambda x:x["fo"])
        lbs=["Lowest Vendor","2nd Lowest","3rd Lowest","4th Lowest","5th Lowest","6th Lowest","7th Lowest","8th Lowest"]
        for si,sd in enumerate(sc_rows):
            W(ws2,row,1,si+1,'A10',s='s'); W(ws2,row,2,lbs[si] if si<len(lbs) else f"{si+1}th",'B10',s='s')
            W(ws2,row,3,sd["sc"]["company"],'C10',s='s')
            for i in range(6):
                t=sd["rt"][i] if i<len(sd["rt"]) else None
                W(ws2,row,4+i,float(t) if t is not None else None,f'{chr(68+i)}10',s='s',nf=NUM)
            W(ws2,row,9,float(sd["fo"]),'I10',s='s',nf=NUM)
            t=sd["ft"]
            for ci,k in enumerate(["adv","mos","prog","ret"],start=10): W(ws2,row,ci,t.get(k,""),f'{chr(64+ci)}10',s='s')
            W(ws2,row,14,t.get("advg",""),'N10',s='s'); W(ws2,row,15,t.get("perfg",""),'O10',s='s')
            W(ws2,row,16,"-",'P10',s='s'); W(ws2,row,17,sd["info0"].get("qref",""),'Q10',s='s')
            ws2.row_dimensions[row].height=ws_sum.row_dimensions.get(10,type('',(),{'height':67.5})()).height or 67.5
            row+=1
    row+=2; W(ws2,row,2,"Recommended Vendor:",'B22',s='s'); row+=1
    W(ws2,row,2,"Back Up Vendor:",'B23',s='s'); row+=2
    W(ws2,row,4,"Budget",'D25',s='s'); W(ws2,row,5,"Selling",'E25',s='s'); row+=1
    W(ws2,row,1,"Allowance",'A26',s='s'); c=W(ws2,row,4,gb,'D26',s='s',nf=NUM); row+=1
    W(ws2,row,1,"Order Value",'A27',s='s'); row+=1
    W(ws2,row,1,"Profit/(Loss)",'A28',s='s'); row+=2
    W(ws2,row,1,"NOTES",'A30',s='s'); W(ws2,row,3,"Specified Material:",'C30',s='s'); row+=1
    W(ws2,row,3,"Warranty as per specs",'C31',s='s')

    buf=io.BytesIO(); wb.save(buf); buf.seek(0)
    return base64.b64encode(buf.read()).decode()
